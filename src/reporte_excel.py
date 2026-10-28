"""Genera el reporte de ventas en Excel con tablas resumen y graficos nativos.

Uso:
    python src/reporte_excel.py            # reporte en reports/
    python src/reporte_excel.py --anio 2017
"""
import argparse
import time
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

RAIZ = Path(__file__).resolve().parents[1]
MODELO = RAIZ / "data" / "model"
REPORTES = RAIZ / "reports"

VERDE = "1E7F5C"
GRIS = "F2F4F3"
FORMATOS = {
    "ventas": '"$"#,##0', "ganancia": '"$"#,##0', "costo": '"$"#,##0',
    "ticket_promedio": '"$"#,##0.00', "margen": "0.0%", "descuento": "0%",
    "tasa_devolucion": "0.0%", "participacion": "0.0%",
}


def cargar_ventas(anio: int | None = None) -> pd.DataFrame:
    """Une la tabla de hechos con sus dimensiones en una sola tabla plana."""
    fact = pd.read_csv(MODELO / "fact_ventas.csv", parse_dates=["fecha_pedido", "fecha_envio"])
    df = (
        fact
        .merge(pd.read_csv(MODELO / "dim_producto.csv"), on="clave_producto")
        .merge(pd.read_csv(MODELO / "dim_geografia.csv", dtype={"codigo_postal": str}), on="clave_geografia")
        .merge(pd.read_csv(MODELO / "dim_cliente.csv"), on="id_cliente")
        .merge(pd.read_csv(MODELO / "dim_region.csv"), on="region", how="left")
    )
    df["anio"] = df["fecha_pedido"].dt.year
    df["anio_mes"] = df["fecha_pedido"].dt.strftime("%Y-%m")
    if anio:
        df = df[df["anio"] == anio]
        if df.empty:
            raise SystemExit(f"No hay ventas para el anio {anio}")
    return df


def resumir(df: pd.DataFrame, por: list[str]) -> pd.DataFrame:
    tabla = (
        df.groupby(por, as_index=False)
        .agg(ventas=("ventas", "sum"), ganancia=("ganancia", "sum"),
             pedidos=("id_pedido", "nunique"), cantidad=("cantidad", "sum"))
    )
    tabla["margen"] = tabla["ganancia"] / tabla["ventas"]
    tabla["participacion"] = tabla["ventas"] / tabla["ventas"].sum()
    return tabla


def indicadores(df: pd.DataFrame) -> list[tuple[str, float, str]]:
    ventas = df["ventas"].sum()
    pedidos = df["id_pedido"].nunique()
    devueltos = df.loc[df["devuelto"] == 1, "id_pedido"].nunique()
    return [
        ("Ventas totales", ventas, FORMATOS["ventas"]),
        ("Ganancia", df["ganancia"].sum(), FORMATOS["ganancia"]),
        ("Margen de ganancia", df["ganancia"].sum() / ventas, FORMATOS["margen"]),
        ("Pedidos", pedidos, "#,##0"),
        ("Clientes", df["id_cliente"].nunique(), "#,##0"),
        ("Ticket promedio", ventas / pedidos, FORMATOS["ticket_promedio"]),
        ("Unidades vendidas", df["cantidad"].sum(), "#,##0"),
        ("Tasa de devolucion", devueltos / pedidos, FORMATOS["tasa_devolucion"]),
        ("Dias promedio de envio", df["dias_envio"].mean(), "0.0"),
    ]


def escribir_tabla(ws, df: pd.DataFrame, fila: int, col: int, nombre: str) -> tuple[int, int]:
    """Escribe df como Tabla de Excel (con filtros) y devuelve la ultima fila y columna."""
    for j, encabezado in enumerate(df.columns):
        ws.cell(fila, col + j, encabezado.replace("_", " ").capitalize())
    for i, registro in enumerate(df.itertuples(index=False), start=1):
        for j, valor in enumerate(registro):
            celda = ws.cell(fila + i, col + j, valor.item() if hasattr(valor, "item") else valor)
            formato = FORMATOS.get(df.columns[j])
            if formato:
                celda.number_format = formato
            elif isinstance(valor, (pd.Timestamp, datetime)):
                celda.number_format = "yyyy-mm-dd"
    ultima_fila, ultima_col = fila + len(df), col + len(df.columns) - 1
    rango = f"{get_column_letter(col)}{fila}:{get_column_letter(ultima_col)}{ultima_fila}"
    tabla = Table(displayName=nombre, ref=rango)
    tabla.tableStyleInfo = TableStyleInfo(name="TableStyleMedium7", showRowStripes=True)
    ws.add_table(tabla)
    return ultima_fila, ultima_col


def ajustar_columnas(ws, minimo: int = 10, maximo: int = 45) -> None:
    for columna in ws.columns:
        largo = max((len(str(c.value)) for c in columna if c.value is not None), default=0)
        ws.column_dimensions[get_column_letter(columna[0].column)].width = max(minimo, min(maximo, largo + 2))


def titulo(ws, texto: str, subtitulo: str) -> None:
    ws["A1"] = texto
    ws["A1"].font = Font(size=16, bold=True, color=VERDE)
    ws["A2"] = subtitulo
    ws["A2"].font = Font(italic=True, color="666666")


def grafico_barras(ws, datos_ref, categorias_ref, titulo_grafico: str, ancla: str,
                   horizontal: bool = False) -> None:
    g = BarChart()
    g.type = "bar" if horizontal else "col"
    g.title = titulo_grafico
    g.style = 10
    g.add_data(datos_ref, titles_from_data=True)
    g.set_categories(categorias_ref)
    g.height, g.width = 8, 16
    if horizontal:
        g.y_axis.scaling.orientation = "minMax"
        g.x_axis.scaling.orientation = "maxMin"
    ws.add_chart(g, ancla)


def hoja_resumen(wb: Workbook, df: pd.DataFrame, periodo: str) -> None:
    ws = wb.active
    ws.title = "Resumen"
    titulo(ws, "Reporte de ventas", f"Periodo: {periodo} | Generado: {datetime.now():%Y-%m-%d %H:%M}")

    borde = Border(bottom=Side(style="thin", color="CCCCCC"))
    for i, (nombre, valor, formato) in enumerate(indicadores(df), start=4):
        ws.cell(i, 1, nombre).font = Font(bold=True)
        celda = ws.cell(i, 2, float(valor))
        celda.number_format = formato
        celda.alignment = Alignment(horizontal="right")
        for c in (1, 2):
            ws.cell(i, c).fill = PatternFill("solid", fgColor=GRIS)
            ws.cell(i, c).border = borde

    por_anio = resumir(df, ["anio"])[["anio", "ventas", "ganancia", "margen", "pedidos"]]
    fila_fin, _ = escribir_tabla(ws, por_anio, 15, 1, "VentasPorAnio")
    grafico_barras(ws, Reference(ws, min_col=2, max_col=3, min_row=15, max_row=fila_fin),
                   Reference(ws, min_col=1, min_row=16, max_row=fila_fin),
                   "Ventas y ganancia por anio", "G4")
    ajustar_columnas(ws)
    ws.column_dimensions["A"].width = 26


def hoja_mensual(wb: Workbook, df: pd.DataFrame) -> None:
    ws = wb.create_sheet("Ventas por mes")
    titulo(ws, "Ventas por mes", "Evolucion mensual de ventas y ganancia")
    tabla = resumir(df, ["anio_mes"])[["anio_mes", "ventas", "ganancia", "margen", "pedidos"]]
    fila_fin, _ = escribir_tabla(ws, tabla, 4, 1, "VentasPorMes")

    g = LineChart()
    g.title = "Ventas y ganancia mensual"
    g.style = 12
    g.add_data(Reference(ws, min_col=2, max_col=3, min_row=4, max_row=fila_fin), titles_from_data=True)
    g.set_categories(Reference(ws, min_col=1, min_row=5, max_row=fila_fin))
    g.height, g.width = 9, 24
    ws.add_chart(g, "G4")
    ajustar_columnas(ws)


def hoja_dimension(wb: Workbook, df: pd.DataFrame, nombre: str, por: list[str],
                   tabla_nombre: str, titulo_grafico: str) -> None:
    ws = wb.create_sheet(nombre)
    titulo(ws, nombre, "Ordenado de mayor a menor venta")
    tabla = resumir(df, por).sort_values("ventas", ascending=False)
    fila_fin, _ = escribir_tabla(ws, tabla, 4, 1, tabla_nombre)
    col_ventas = len(por) + 1
    grafico_barras(ws, Reference(ws, min_col=col_ventas, min_row=4, max_row=fila_fin),
                   Reference(ws, min_col=len(por), min_row=5, max_row=fila_fin),
                   titulo_grafico, f"{get_column_letter(len(tabla.columns) + 2)}4",
                   horizontal=len(tabla) > 6)
    ajustar_columnas(ws)


def hoja_top_productos(wb: Workbook, df: pd.DataFrame, n: int = 10) -> None:
    ws = wb.create_sheet("Top productos")
    titulo(ws, f"Top {n} productos", "Por ventas, con su ganancia y margen")
    tabla = (
        resumir(df, ["producto", "categoria"])
        .nlargest(n, "ventas")[["producto", "categoria", "ventas", "ganancia", "margen", "cantidad"]]
    )
    fila_fin, _ = escribir_tabla(ws, tabla, 4, 1, "TopProductos")
    grafico_barras(ws, Reference(ws, min_col=3, min_row=4, max_row=fila_fin),
                   Reference(ws, min_col=1, min_row=5, max_row=fila_fin),
                   f"Top {n} productos por ventas", "H4", horizontal=True)
    ajustar_columnas(ws)


def hoja_matriz(wb: Workbook, df: pd.DataFrame) -> None:
    """Cruce region x categoria, equivalente a una tabla dinamica."""
    ws = wb.create_sheet("Region x Categoria")
    titulo(ws, "Ventas por region y categoria", "Ventas (arriba) y margen (abajo)")
    ventas = df.pivot_table(index="region", columns="categoria", values="ventas",
                            aggfunc="sum", margins=True, margins_name="Total")
    ganancia = df.pivot_table(index="region", columns="categoria", values="ganancia",
                              aggfunc="sum", margins=True, margins_name="Total")
    bloques = [(ventas.reset_index(), "MatrizVentas", '"$"#,##0'),
               ((ganancia / ventas).reset_index(), "MatrizMargen", "0.0%")]
    fila = 4
    for tabla, nombre, formato in bloques:
        tabla.columns = [str(c) for c in tabla.columns]
        fila_fin, col_fin = escribir_tabla(ws, tabla, fila, 1, nombre)
        for r in ws.iter_rows(min_row=fila + 1, max_row=fila_fin, min_col=2, max_col=col_fin):
            for c in r:
                c.number_format = formato
        fila = fila_fin + 3
    ajustar_columnas(ws)


def hoja_datos(wb: Workbook, df: pd.DataFrame) -> None:
    ws = wb.create_sheet("Datos")
    columnas = ["id_pedido", "fecha_pedido", "region", "estado", "ciudad", "segmento",
                "categoria", "subcategoria", "producto", "ventas", "cantidad",
                "descuento", "ganancia", "devuelto"]
    escribir_tabla(ws, df[columnas].sort_values("fecha_pedido"), 1, 1, "DatosVentas")
    ajustar_columnas(ws)
    ws.freeze_panes = "A2"


def generar(anio: int | None = None, destino: Path = REPORTES) -> tuple[Path, float]:
    inicio = time.perf_counter()
    df = cargar_ventas(anio)
    periodo = str(anio) if anio else f"{df['anio'].min()} - {df['anio'].max()}"

    wb = Workbook()
    hoja_resumen(wb, df, periodo)
    hoja_mensual(wb, df)
    hoja_dimension(wb, df, "Por region", ["region", "responsable"], "VentasRegion", "Ventas por region")
    hoja_dimension(wb, df, "Por categoria", ["categoria", "subcategoria"], "VentasCategoria",
                   "Ventas por subcategoria")
    hoja_top_productos(wb, df)
    hoja_matriz(wb, df)
    hoja_datos(wb, df)

    destino.mkdir(parents=True, exist_ok=True)
    sufijo = anio or "completo"
    ruta = destino / f"reporte_ventas_{sufijo}_{datetime.now():%Y%m%d}.xlsx"
    wb.save(ruta)
    return ruta, time.perf_counter() - inicio


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--anio", type=int, help="Filtra el reporte a un solo anio")
    args = parser.parse_args()
    ruta, segundos = generar(args.anio)
    print(f"Reporte generado: {ruta}")
    print(f"Tiempo: {segundos:.2f} s")
