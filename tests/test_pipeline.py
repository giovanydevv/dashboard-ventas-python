import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

import limpieza  # noqa: E402
import modelo_datos  # noqa: E402

CABECERA = ",".join(limpieza.COLUMNAS_PEDIDOS)
FILA = "{id},CA-1,11/8/2017,11/11/2017,Second Class,C-1,Ana,Consumer,United States,{ciudad},{estado},{cp},South,P-1,Furniture,Chairs,Silla,100.5,2,0.2,{ganancia}"


@pytest.fixture
def csv_sucio(tmp_path: Path) -> Path:
    filas = [
        CABECERA,
        FILA.format(id=1, ciudad="Henderson", estado="Kentucky", cp="42420", ganancia="10"),
        FILA.format(id=2, ciudad="Henderson", estado="Kentucky", cp="42420", ganancia="10"),  # duplicado
        FILA.format(id=3, ciudad=" Burlington ", estado="Vermont", cp="", ganancia="-5"),       # sin CP y espacios
        FILA.format(id=4, ciudad="Boston", estado="Massachusetts", cp="2108", ganancia="7"),    # CP sin cero
        "Person,Region",
        "Ana Lopez,South",
        "Returned,Order ID",
        "Yes,CA-1",
        "Yes,CA-1",
    ]
    ruta = tmp_path / "raw.csv"
    ruta.write_text("\n".join(filas), encoding="utf-8")
    return ruta


def test_separa_las_tres_tablas(csv_sucio):
    secciones = limpieza.separar_secciones(csv_sucio)
    assert list(secciones) == ["pedidos", "responsables", "devoluciones"]
    assert len(secciones["pedidos"]) == 4


def test_limpieza_corrige_duplicados_nulos_y_formatos(csv_sucio, tmp_path):
    registro = limpieza.ejecutar(csv_sucio, tmp_path / "out")
    df = pd.read_csv(tmp_path / "out" / "pedidos_limpios.csv", dtype={"codigo_postal": str})

    assert registro["duplicados_eliminados"] == 1
    assert len(df) == 3
    assert set(df["codigo_postal"]) == {"42420", "05401", "02108"}
    assert "Burlington" in set(df["ciudad"])
    assert df["fecha_pedido"].iloc[0] == "2017-11-08"
    assert registro["devoluciones_pedidos_unicos"] == 1


def test_modelo_no_pierde_filas_y_marca_devoluciones(csv_sucio, tmp_path, monkeypatch):
    limpieza.ejecutar(csv_sucio, tmp_path / "out")
    monkeypatch.setattr(modelo_datos, "PROCESADOS", tmp_path / "out")
    tablas = modelo_datos.construir(*modelo_datos.cargar())

    fact = tablas["fact_ventas"]
    assert len(fact) == 3
    assert fact["devuelto"].eq(1).all()
    assert fact["costo"].tolist() == pytest.approx([90.5, 105.5, 93.5])
    assert tablas["dim_fecha"]["fecha"].min() == pd.Timestamp("2017-01-01")


def test_reporte_real_tiene_todas_las_hojas(tmp_path):
    import reporte_excel

    if not (reporte_excel.MODELO / "fact_ventas.csv").exists():
        pytest.skip("Ejecuta primero python main.py")
    ruta, _ = reporte_excel.generar(anio=2017, destino=tmp_path)
    hojas = load_workbook(ruta).sheetnames
    assert hojas == ["Resumen", "Ventas por mes", "Por region", "Por categoria",
                     "Top productos", "Region x Categoria", "Datos"]
